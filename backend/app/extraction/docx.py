from __future__ import annotations

import io

import docx

from app.extraction.textutil import clean_text
from app.extraction.types import ExtractedText
from app.jobs.errors import PermanentJobError


def extract_docx(data: bytes) -> ExtractedText:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise PermanentJobError("corrupt_docx") from exc
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return ExtractedText(pages=[clean_text("\n".join(parts))], source="docx")
