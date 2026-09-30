"""Spreadsheet import and the quarterly hiring export.

Cells are data. They are never instructions, and they are not copied into the audit log.
"""

from __future__ import annotations

import re
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.domain.actor import Actor
from app.domain.confidence import Weights, assess
from app.domain.identity import name_similarity
from app.extraction.contacts import normalize_email, normalize_phone
from app.models import (
    Application,
    Candidate,
    CandidateContact,
    CandidateMatch,
    ExcelImport,
    ExcelImportRow,
    HiringRequest,
    RequestRequirement,
    SharingRecord,
    User,
)
from app.services.audit import write_audit
from app.services.compare import contact_pairs
from app.services.reviews import open_excel_relationship_review

MAX_ROWS = 5000
NAME_RE = re.compile(r"^[\w][\w .'\-]{0,120}$", re.UNICODE)
INSTRUCTION_RE = re.compile(
    r"(ignore\s+(all\s+)?previous|system\s*:|you\s+are\s+now|<\s*script|drop\s+table|union\s+select)",
    re.IGNORECASE,
)
HEADER_ALIASES = {
    "name": "name",
    "full name": "name",
    "candidate": "name",
    "candidate name": "name",
    "email": "email",
    "e-mail": "email",
    "email address": "email",
    "phone": "phone",
    "mobile": "phone",
    "telephone": "phone",
    "phone number": "phone",
}
EXPORT_HEADERS = (
    "request_id",
    "company",
    "job_title",
    "status",
    "received_on",
    "owner",
    "requirements",
    "candidates_found",
    "candidates_shared",
    "candidate_ids",
)


class ExcelError(Exception):
    def __init__(self, code: str) -> None:
        self.code = code
        super().__init__(code)


def import_candidates(session: Session, *, actor: Actor, filename: str | None, data: bytes) -> ExcelImport:
    if actor.kind != "user" or actor.user_id is None:
        raise ExcelError("recruiter_required")
    if not filename or not filename.lower().endswith(".xlsx"):
        raise ExcelError("unsupported_file")
    if len(data) > get_settings().max_upload_bytes:
        raise ExcelError("file_too_large")
    sheet_rows, sheet_count, ignored_columns = _read_sheet(data)
    batch = ExcelImport(
        actor_user_id=actor.user_id,
        sheet_count=sheet_count,
        row_count=0,
        valid_count=0,
        invalid_count=0,
        duplicate_count=0,
        relationship_count=0,
        created_count=0,
        ignored_column_count=ignored_columns,
    )
    session.add(batch)
    session.flush()
    seen: list[tuple[int, set[str], str]] = []
    for row_number, cells in sheet_rows:
        outcome = _import_row(session, batch, row_number, cells, seen, actor)
        if outcome == "created":
            batch.created_count += 1
            batch.valid_count += 1
        elif outcome == "existing_relationship":
            batch.relationship_count += 1
            batch.valid_count += 1
        elif outcome == "duplicate_row":
            batch.duplicate_count += 1
            batch.valid_count += 1
        else:
            batch.invalid_count += 1
        batch.row_count += 1
    write_audit(
        session,
        actor=actor,
        action="excel_import_completed",
        entity_type="excel_import",
        entity_id=batch.id,
        detail={
            "count": batch.row_count,
            "created_count": batch.created_count,
            "invalid_count": batch.invalid_count,
            "duplicate_count": batch.duplicate_count,
            "relationship_count": batch.relationship_count,
            "result": "completed",
        },
    )
    session.flush()
    return batch


def export_hiring_activity(session: Session, *, actor: Actor, start: date, end: date) -> bytes:
    if actor.kind != "user" or actor.user_id is None:
        raise ExcelError("recruiter_required")
    if end < start:
        raise ExcelError("invalid_range")
    start_at = datetime(start.year, start.month, start.day, tzinfo=UTC)
    end_at = datetime(end.year, end.month, end.day, tzinfo=UTC) + timedelta(days=1)
    requests = session.scalars(
        select(HiringRequest)
        .where(HiringRequest.received_at >= start_at, HiringRequest.received_at < end_at)
        .order_by(HiringRequest.received_at)
    ).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hiring activity"
    sheet.append(list(EXPORT_HEADERS))
    for request in requests:
        sheet.append(_export_row(session, request))
    write_audit(
        session,
        actor=actor,
        action="hiring_activity_exported",
        entity_type="hiring_request",
        entity_id=None,
        detail={
            "count": len(requests),
            "period_start": start.isoformat(),
            "period_end": end.isoformat(),
            "result": "completed",
        },
    )
    session.flush()
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _read_sheet(data: bytes) -> tuple[list[tuple[int, dict[str, str]]], int, int]:
    try:
        workbook = load_workbook(BytesIO(data), data_only=False, read_only=True)
    except Exception as exc:
        raise ExcelError("invalid_workbook") from exc
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise ExcelError("empty_workbook") from exc
        columns, ignored = _columns(header)
        if "name" not in columns:
            raise ExcelError("missing_name_column")
        parsed: list[tuple[int, dict[str, str]]] = []
        for offset, raw in enumerate(rows, start=2):
            if len(parsed) >= MAX_ROWS:
                raise ExcelError("too_many_rows")
            if raw is None or all(_plain(value) == "" for value in raw):
                continue
            cells: dict[str, str] = {}
            unsafe = False
            for field, index in columns.items():
                value = raw[index] if index < len(raw) else None
                if _is_formula(value):
                    unsafe = True
                    cells[field] = ""
                    continue
                cells[field] = _plain(value)
            if unsafe:
                cells["__unsafe"] = "1"
            parsed.append((offset, cells))
        return parsed, len(workbook.worksheets), ignored
    finally:
        workbook.close()


def _columns(header: tuple | list) -> tuple[dict[str, int], int]:
    columns: dict[str, int] = {}
    ignored = 0
    for index, value in enumerate(header):
        label = _plain(value).casefold()
        if not label:
            continue
        field = HEADER_ALIASES.get(label)
        if field is None or field in columns:
            ignored += 1
            continue
        columns[field] = index
    return columns, ignored


def _import_row(
    session: Session,
    batch: ExcelImport,
    row_number: int,
    cells: dict[str, str],
    seen: list[tuple[int, set[str], str]],
    actor: Actor,
) -> str:
    errors = _row_errors(cells)
    if errors:
        _add_row(session, batch.id, row_number, "invalid", errors)
        return "invalid"
    name = cells["name"].strip()
    email = normalize_email(cells.get("email", "")) if cells.get("email") else None
    phone = normalize_phone(cells.get("phone", "")) if cells.get("phone") else None
    keys = _keys(name, email, phone)
    prior = _prior_row(seen, keys, name)
    if prior is not None:
        _add_row(session, batch.id, row_number, "duplicate_row", [], duplicate_of_row=prior)
        return "duplicate_row"
    emails = [(email, email)] if email else []
    phones = [(phone, phone)] if phone else []
    assessment = assess(
        contact_pairs(session, emails=emails, phones=phones, name=name),
        weights=_weights(),
        high_confidence=get_settings().identity_high_confidence,
        relationship_cutoff=get_settings().identity_relationship_cutoff,
    )
    if assessment.comparable:
        row = _add_row(
            session,
            batch.id,
            row_number,
            "existing_relationship",
            [],
            candidate_id=assessment.candidate_id,
            relationship=assessment.relationship,
            proposed_name=name[:200],
            proposed_email=email,
            proposed_phone=phone,
        )
        open_excel_relationship_review(session, row, assessment, actor)
        seen.append((row_number, keys, name))
        return "existing_relationship"
    candidate = Candidate(source="excel", display_name=name[:200])
    session.add(candidate)
    session.flush()
    received = datetime.now(UTC)
    session.add(
        Application(
            candidate_id=candidate.id,
            document_id=None,
            source="excel",
            received_at=received,
        )
    )
    if email:
        _add_contact(session, candidate.id, "email", email)
    if phone:
        _add_contact(session, candidate.id, "phone", phone)
    _add_row(session, batch.id, row_number, "created", [], candidate_id=candidate.id)
    seen.append((row_number, keys, name))
    return "created"


def _row_errors(cells: dict[str, str]) -> list[str]:
    errors: list[str] = []
    if cells.get("__unsafe"):
        errors.append("untrusted_content")
    name = cells.get("name", "").strip()
    email = cells.get("email", "").strip()
    phone = cells.get("phone", "").strip()
    if not name:
        errors.append("missing_name")
    elif not _safe_name(name):
        errors.append("untrusted_content")
    if email and normalize_email(email) is None:
        errors.append("malformed_email")
    if phone and normalize_phone(phone) is None:
        errors.append("malformed_phone")
    if _instruction(email) or _instruction(phone):
        errors.append("untrusted_content")
    return list(dict.fromkeys(errors))


def _safe_name(value: str) -> bool:
    if _instruction(value):
        return False
    if "@" in value or any(ord(char) < 32 for char in value):
        return False
    return NAME_RE.fullmatch(value) is not None


def _instruction(value: str) -> bool:
    return bool(value and INSTRUCTION_RE.search(value))


def _keys(name: str, email: str | None, phone: str | None) -> set[str]:
    keys: set[str] = set()
    if email:
        keys.add(f"email:{email}")
    if phone:
        keys.add(f"phone:{phone}")
    return keys


def _prior_row(seen: list[tuple[int, set[str], str]], keys: set[str], name: str) -> int | None:
    for row_number, prior_keys, prior_name in seen:
        if keys and prior_keys and keys & prior_keys:
            return row_number
        if not keys and not prior_keys and name_similarity(name, prior_name) == 1.0:
            return row_number
    return None


def _add_contact(session: Session, candidate_id: UUID, kind: str, normalized: str) -> None:
    session.add(
        CandidateContact(
            candidate_id=candidate_id,
            kind=kind,
            value_raw=normalized[:320],
            value_normalized=normalized[:320],
            is_primary=True,
        )
    )


def _add_row(
    session: Session,
    import_id: UUID,
    row_number: int,
    outcome: str,
    errors: list[str],
    *,
    candidate_id: UUID | None = None,
    relationship: str | None = None,
    duplicate_of_row: int | None = None,
    proposed_name: str | None = None,
    proposed_email: str | None = None,
    proposed_phone: str | None = None,
) -> ExcelImportRow:
    row = ExcelImportRow(
        import_id=import_id,
        row_number=row_number,
        outcome=outcome,
        error_codes=errors,
        candidate_id=candidate_id,
        relationship=relationship,
        duplicate_of_row=duplicate_of_row,
        proposed_name=proposed_name,
        proposed_email=proposed_email,
        proposed_phone=proposed_phone,
    )
    session.add(row)
    session.flush()
    return row


def _weights() -> Weights:
    settings = get_settings()
    return Weights(
        email=settings.identity_weight_email,
        phone=settings.identity_weight_phone,
        name=settings.identity_weight_name,
        dates=settings.identity_weight_dates,
        explicit_version=settings.identity_weight_explicit_version,
        employment=settings.identity_weight_employment,
        education=settings.identity_weight_education,
        skills=settings.identity_weight_skills,
        semantic=settings.identity_weight_semantic,
    )


def _export_row(session: Session, request: HiringRequest) -> list[object]:
    owner = session.get(User, request.owner_user_id) if request.owner_user_id else None
    labels = session.scalars(
        select(RequestRequirement.label).where(
            RequestRequirement.request_id == request.id,
            RequestRequirement.review_status == "approved",
        )
    ).all()
    match_ids = session.scalars(
        select(CandidateMatch.candidate_id).where(CandidateMatch.request_id == request.id)
    ).all()
    shared_ids = session.scalars(
        select(SharingRecord.candidate_id).where(SharingRecord.request_id == request.id)
    ).all()
    found = session.scalar(
        select(func.count()).select_from(CandidateMatch).where(CandidateMatch.request_id == request.id)
    )
    shared = session.scalar(
        select(func.count()).select_from(SharingRecord).where(SharingRecord.request_id == request.id)
    )
    candidate_ids = sorted({str(item) for item in (*match_ids, *shared_ids)})
    return [
        str(request.id),
        _public_text(request.company_name),
        _public_text(request.job_title),
        request.status,
        request.received_at.date().isoformat(),
        _public_text(owner.full_name) if owner else "",
        _public_text(", ".join(label for label in labels if _public_text(label))),
        int(found or 0),
        int(shared or 0),
        ",".join(candidate_ids),
    ]


def _public_text(value: str | None) -> str:
    if not value:
        return ""
    if "@" in value or re.search(r"\d{8,}", value):
        return ""
    return value[:300]


def _plain(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value:
            return ""
        if value.is_integer():
            return str(int(value))
        return ""
    if isinstance(value, datetime | date):
        return ""
    return str(value).replace("\x00", "").strip()


def _is_formula(value: object) -> bool:
    return isinstance(value, str) and value.strip().startswith("=")
