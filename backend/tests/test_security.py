"""Authentication, audit immutability, and PII boundaries."""

import os
from datetime import UTC, datetime
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import func, select, text
from sqlalchemy.exc import DBAPIError

from app.bootstrap import provision
from app.config import Settings, get_settings, reset_settings
from app.domain.actor import SYSTEM, Actor
from app.domain.transitions import TransitionError
from app.main import app
from app.models import (
    AuditEvent,
    Candidate,
    CandidateContact,
    CandidateDocument,
    DocumentChunk,
    ExtractionEvidence,
    HiringRequest,
    Shortlist,
    User,
)
from app.services.audit import write_audit
from app.services.excel_ops import ExcelError, export_hiring_activity, import_candidates
from app.services.hiring import shortlist_match
from app.services.merge import MergeError, merge_profiles
from app.services.sharing import send_share
from tests.test_mail import ENGINEER, RECRUITER

CONTACT = "ada@example.com"
QUOTE = "Reconcile ledger accounts"
MISSING = "00000000-0000-0000-0000-000000000099"
PUBLIC_PATHS = {"/health", "/openapi.json", "/docs", "/redoc"}


class _SilentMail:
    def send(self, message) -> None:
        raise AssertionError("mail was sent")


def _restore_local_settings() -> None:
    os.environ["ENVIRONMENT"] = "local"
    os.environ["HUB_RECRUITER_TOKEN"] = "dev-recruiter-token"
    os.environ["HUB_ENGINEER_TOKEN"] = "dev-engineer-token"
    reset_settings()
    get_settings()


def _ready(session):
    provision(session)
    session.commit()
    return session.scalar(select(User).where(User.role == "recruiter"))


def _candidate_with_pii(session):
    recruiter = _ready(session)
    now = datetime.now(UTC)
    candidate = Candidate(source="upload", display_name="Ada Lovelace", summary="Ledger specialist")
    session.add(candidate)
    session.flush()
    document = CandidateDocument(
        candidate_id=candidate.id,
        sha256="a" * 64,
        original_filename="ada.pdf",
        media_type="application/pdf",
        byte_size=32,
        storage_key="security/ada.pdf",
        status="profiled",
        is_current=True,
        processing_version=1,
        received_at=now,
    )
    session.add(document)
    session.flush()
    session.add(
        CandidateContact(
            candidate_id=candidate.id,
            kind="email",
            value_raw=CONTACT,
            value_normalized=CONTACT,
            is_primary=True,
        )
    )
    chunk = DocumentChunk(
        document_id=document.id,
        chunk_index=0,
        content=f"Email {CONTACT}. {QUOTE}",
        content_hash="b" * 64,
    )
    session.add(chunk)
    session.flush()
    session.add(
        ExtractionEvidence(
            document_id=document.id,
            candidate_id=candidate.id,
            chunk_id=chunk.id,
            entity_type="contact",
            field_name="email",
            extracted_value=CONTACT,
            source_quote=CONTACT,
            quote_verified=True,
        )
    )
    session.add(
        ExtractionEvidence(
            document_id=document.id,
            candidate_id=candidate.id,
            chunk_id=chunk.id,
            entity_type="experience",
            field_name="description",
            extracted_value="Engineer",
            source_quote=QUOTE,
            quote_verified=True,
        )
    )
    session.commit()
    return recruiter, candidate, document


def test_tokens_are_required_settings_and_demo_tokens_stay_local(monkeypatch):
    assert Settings.model_fields["hub_recruiter_token"].is_required()
    assert Settings.model_fields["hub_engineer_token"].is_required()
    monkeypatch.delenv("HUB_RECRUITER_TOKEN", raising=False)
    monkeypatch.delenv("HUB_ENGINEER_TOKEN", raising=False)
    monkeypatch.setenv("ENVIRONMENT", "production")
    reset_settings()
    try:
        with pytest_raises_settings():
            Settings(_env_file=None)
    finally:
        _restore_local_settings()

    for environment in ("production", "staging"):
        monkeypatch.setenv("ENVIRONMENT", environment)
        monkeypatch.setenv("HUB_RECRUITER_TOKEN", "dev-recruiter-token")
        monkeypatch.setenv("HUB_ENGINEER_TOKEN", "dev-engineer-token")
        reset_settings()
        try:
            with pytest_raises_runtime():
                get_settings()
        finally:
            _restore_local_settings()

    monkeypatch.setenv("ENVIRONMENT", "production")
    monkeypatch.setenv("HUB_RECRUITER_TOKEN", "recruiter-production-token-value")
    monkeypatch.setenv("HUB_ENGINEER_TOKEN", "engineer-production-token-value")
    reset_settings()
    try:
        loaded = get_settings()
        assert loaded.environment == "production"
        assert loaded.hub_recruiter_token != loaded.hub_engineer_token
    finally:
        _restore_local_settings()


def pytest_raises_settings():
    import pytest

    return pytest.raises(ValidationError)


def pytest_raises_runtime():
    import pytest

    return pytest.raises(RuntimeError, match="development tokens")


def test_missing_token_is_401_on_every_protected_route():
    with TestClient(app) as client:
        assert client.get("/health").status_code == 200
        schema = client.get("/openapi.json").json()
        checked = 0
        for path, operations in schema["paths"].items():
            if path in PUBLIC_PATHS or path.startswith("/docs"):
                continue
            for method in operations:
                if method not in {"get", "post", "patch", "put", "delete"}:
                    continue
                url = path
                for name in _path_names(path):
                    url = url.replace("{" + name + "}", MISSING)
                response = client.request(method.upper(), url)
                assert response.status_code == 401, (method, path, response.status_code)
                assert response.json()["detail"] == "unauthorized"
                checked += 1
        assert checked >= 40


def test_wrong_role_is_forbidden(session):
    _ready(session)
    with TestClient(app) as client:
        recruiter_routes = (
            ("get", "/dashboard"),
            ("get", "/work-queue"),
            ("get", "/search"),
            ("get", "/hiring-requests"),
            ("get", "/candidates/" + MISSING),
            ("get", "/documents/" + MISSING + "/text"),
        )
        for method, path in recruiter_routes:
            denied = client.request(method.upper(), path, headers=ENGINEER)
            assert denied.status_code == 403, path
            assert denied.json()["detail"] == "recruiter_required"
            assert CONTACT not in denied.text
        created = client.post(
            "/problem-reports",
            headers=RECRUITER,
            json={
                "category": "intake",
                "title": "Upload stayed on reading",
                "description": "A PDF remained queued.",
                "entity_type": "none",
                "context": {"route": "/source"},
            },
        )
        assert created.status_code == 201
        report_id = created.json()["id"]
        engineer_view = client.get("/problem-reports", headers=ENGINEER)
        assert engineer_view.status_code == 200
        assert engineer_view.json()["reports"][0]["id"] == report_id
        assert CONTACT not in engineer_view.text
        recruiter_update = client.patch(
            f"/problem-reports/{report_id}",
            headers=RECRUITER,
            json={"status": "resolved", "note": "Closed from the wrong role."},
        )
        assert recruiter_update.status_code == 403
        assert recruiter_update.json()["detail"] == "engineer_required"


def test_audit_update_is_rejected_and_inserts_still_work(session):
    recruiter = _ready(session)
    event = write_audit(
        session,
        actor=Actor(kind="user", user_id=recruiter.id),
        action="security_probe",
        entity_type="candidate",
        entity_id=None,
        detail={"result": "recorded"},
    )
    session.commit()
    event_id = event.id
    row = session.get(AuditEvent, event_id)
    row.action = "tampered"
    with _dbapi():
        session.commit()
    session.rollback()
    kept = session.get(AuditEvent, event_id)
    assert kept.action == "security_probe"
    assert kept.detail == {"result": "recorded"}
    write_audit(
        session,
        actor=Actor(kind="user", user_id=recruiter.id),
        action="security_probe_again",
        entity_type="candidate",
        entity_id=None,
        detail={},
    )
    session.commit()
    assert session.scalar(select(AuditEvent).where(AuditEvent.action == "security_probe_again")) is not None


def test_audit_delete_is_rejected(session):
    recruiter = _ready(session)
    event = write_audit(
        session,
        actor=Actor(kind="user", user_id=recruiter.id),
        action="security_probe",
        entity_type="candidate",
        entity_id=None,
        detail={"count": 1},
    )
    session.commit()
    event_id = event.id
    with _dbapi():
        session.execute(text("DELETE FROM audit_events WHERE id = :id"), {"id": event_id})
    session.rollback()
    session.delete(session.get(AuditEvent, event_id))
    with _dbapi():
        session.commit()
    session.rollback()
    assert session.get(AuditEvent, event_id).action == "security_probe"
    assert session.scalar(select(func.count()).select_from(AuditEvent)) == 1


def test_candidate_pii_stays_with_the_recruiter(session):
    _recruiter, candidate, document = _candidate_with_pii(session)
    with TestClient(app) as client:
        visible = client.get(f"/candidates/{candidate.id}", headers=RECRUITER)
        hidden = client.get(f"/candidates/{candidate.id}", headers=ENGINEER)
        text = client.get(f"/documents/{document.id}/text", headers=RECRUITER)
        hidden_text = client.get(f"/documents/{document.id}/text", headers=ENGINEER)
        document_view = client.get(f"/documents/{document.id}", headers=RECRUITER)
        hidden_document = client.get(f"/documents/{document.id}", headers=ENGINEER)
        listing = client.get("/documents", headers=ENGINEER)
        activity = client.get(f"/documents/{document.id}/activity", headers=ENGINEER)
        reports = client.get("/problem-reports", headers=ENGINEER)
    assert visible.status_code == 200
    assert visible.json()["contacts"][0]["value"] == CONTACT
    assert any(item["quote"] == QUOTE for item in visible.json()["evidence"])
    assert hidden.status_code == 403
    assert CONTACT not in hidden.text
    assert QUOTE not in hidden.text
    assert text.status_code == 200
    assert CONTACT in text.json()["text"]
    assert hidden_text.status_code == 403
    assert CONTACT not in hidden_text.text
    assert document_view.status_code == 200
    assert document_view.json()["extracted_contacts"][0]["value"] == CONTACT
    assert hidden_document.status_code == 403
    assert CONTACT not in hidden_document.text
    assert listing.status_code == 200
    assert CONTACT not in listing.text
    assert "ada.pdf" not in listing.text
    assert activity.status_code == 200
    assert CONTACT not in activity.text
    assert QUOTE not in activity.text
    assert reports.status_code == 200
    session.expire_all()
    audits = session.scalars(select(AuditEvent)).all()
    assert any(row.action == "pii_text_access" for row in audits)
    assert all(CONTACT not in str(row.detail) for row in audits)
    assert all("@" not in str(row.detail) for row in audits)


def test_sharing_shortlist_and_merge_stay_recruiter_actions(session):
    recruiter = _ready(session)
    left = Candidate(source="upload", display_name="Ada Lovelace")
    right = Candidate(source="upload", display_name="Grace Hopper")
    session.add_all([left, right])
    session.commit()
    with _merge_error("human_required"):
        merge_profiles(session, left.id, right.id, SYSTEM, confirm=True)
    with _transition("human_required"):
        shortlist_match(session, left.id, right.id, SYSTEM)
    with _transition("human_required"):
        send_share(session, left.id, right.id, SYSTEM, _SilentMail())
    engineer = session.scalar(select(User).where(User.role == "engineer"))
    with _transition("forbidden"):
        shortlist_match(session, left.id, right.id, Actor(kind="user", user_id=engineer.id))
    with TestClient(app) as client:
        for path in (
            f"/hiring-requests/{MISSING}/matches/{MISSING}/shortlist",
            f"/hiring-requests/{MISSING}/candidates/{MISSING}/share",
            f"/hiring-requests/{MISSING}/candidates/{MISSING}/share/approve",
        ):
            missing = client.post(path)
            denied = client.post(path, headers=ENGINEER)
            assert missing.status_code == 401, path
            assert denied.status_code == 403, path
            assert denied.json()["detail"] == "recruiter_required"
        merge_missing = client.post(
            "/profile-merges",
            json={"left_id": str(left.id), "right_id": str(right.id), "confirm": True},
        )
        merge_denied = client.post(
            "/profile-merges",
            headers=ENGINEER,
            json={"left_id": str(left.id), "right_id": str(right.id), "confirm": True},
        )
        assert merge_missing.status_code == 401
        assert merge_denied.status_code == 403
        still_recruiter = client.post(
            f"/hiring-requests/{MISSING}/matches/{MISSING}/shortlist",
            headers=RECRUITER,
        )
        assert still_recruiter.status_code == 404
    session.expire_all()
    assert session.get(Candidate, left.id).merged_into_id is None
    assert session.scalar(select(func.count()).select_from(Shortlist)) == 0
    assert session.scalar(select(AuditEvent).where(AuditEvent.action == "profiles_merged")) is None
    assert recruiter.role == "recruiter"


def test_excel_import_and_export_stay_recruiter_actions(session):
    recruiter = _ready(session)
    session.add(
        HiringRequest(
            company_name="Northwind",
            job_title="Analyst",
            description=f"Reach {CONTACT} about the role",
            status="open",
            received_at=datetime(2026, 2, 1, tzinfo=UTC),
            owner_user_id=recruiter.id,
        )
    )
    session.commit()
    with _excel("recruiter_required"):
        import_candidates(session, actor=SYSTEM, filename="people.xlsx", data=b"not-a-sheet")
    with _excel("recruiter_required"):
        export_hiring_activity(session, actor=SYSTEM, start=datetime(2026, 1, 1).date(), end=datetime(2026, 3, 31).date())
    with TestClient(app) as client:
        missing_import = client.post(
            "/excel/candidates",
            files={"file": ("people.xlsx", b"not-a-sheet", "application/vnd.ms-excel")},
        )
        denied_import = client.post(
            "/excel/candidates",
            headers=ENGINEER,
            files={"file": ("people.xlsx", b"not-a-sheet", "application/vnd.ms-excel")},
        )
        missing_export = client.get(
            "/hiring-activity/export",
            params={"start": "2026-01-01", "end": "2026-03-31"},
        )
        denied_export = client.get(
            "/hiring-activity/export",
            headers=ENGINEER,
            params={"start": "2026-01-01", "end": "2026-03-31"},
        )
        exported = client.get(
            "/hiring-activity/export",
            headers=RECRUITER,
            params={"start": "2026-01-01", "end": "2026-03-31"},
        )
    assert missing_import.status_code == 401
    assert denied_import.status_code == 403
    assert denied_import.json()["detail"] == "recruiter_required"
    assert CONTACT not in denied_import.text
    assert missing_export.status_code == 401
    assert denied_export.status_code == 403
    assert CONTACT not in denied_export.text
    assert exported.status_code == 200
    assert CONTACT not in exported.content.decode("latin1", errors="ignore")
    book = load_workbook(BytesIO(exported.content))
    cells = [str(cell) for row in book.active.iter_rows(values_only=True) for cell in row]
    assert CONTACT not in " ".join(cells)
    session.expire_all()
    event = session.scalar(select(AuditEvent).where(AuditEvent.action == "hiring_activity_exported"))
    assert event is not None
    assert "@" not in str(event.detail)
    assert session.scalar(select(func.count()).select_from(Candidate)) == 0


def _path_names(path: str) -> list[str]:
    names = []
    for part in path.split("/"):
        if part.startswith("{") and part.endswith("}"):
            names.append(part[1:-1])
    return names


def _dbapi():
    import pytest

    return pytest.raises(DBAPIError, match="append-only")


def _merge_error(code: str):
    import pytest

    return pytest.raises(MergeError, match=code)


def _transition(code: str):
    import pytest

    return pytest.raises(TransitionError, match=code)


def _excel(code: str):
    import pytest

    return pytest.raises(ExcelError, match=code)
