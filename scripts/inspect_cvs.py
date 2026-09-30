"""Inspect the local cvs/ folder without sending any file contents out.

Reports file types, whether a local extractor can read text, and filename
patterns that look like duplicates or later versions. Extracted text is
counted and then discarded.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from pathlib import Path

import docx
import pymupdf

# A short CV still has far more than this. Below it, a PDF is treated as a
# scan or a blank page rather than a usable text layer.
MIN_ALNUM = 80

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".tif", ".tiff", ".bmp", ".webp"}


def token(word: str) -> re.Pattern[str]:
    """Match a word even when the filename uses underscores instead of spaces."""
    return re.compile(rf"(?<![A-Za-z0-9]){word}(?![A-Za-z0-9])", re.IGNORECASE)


# Filename tokens that usually mark a later save of the same document.
# (2026) is a year, not a copy number. (1) and (16) are copy numbers.
VERSION_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    ("copy index", re.compile(r"\((?!(?:19|20)\d{2}\))\d+\)")),
    ("copy", token("copy")),
    ("final", token("final")),
    ("draft", token("draft")),
    ("latest", token("latest")),
    ("updated", token("updated")),
    ("update", token("update")),
    ("revised", token("revised")),
    ("official", token("official")),
    ("master", token("master")),
    ("improved", token("improved")),
    ("version tag", re.compile(r"(?<![A-Za-z0-9])(?:v\d+|vf|version)(?![A-Za-z0-9])", re.IGNORECASE)),
    ("updated (hebrew)", re.compile("\u05de\u05e2\u05d5\u05d3\u05db\u05df")),
]

# Filename words for files that are probably not a candidate CV.
NON_CV_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    ("invoice", token("invoice")),
    ("job description", re.compile(r"(?<![A-Za-z0-9])job description(?![A-Za-z0-9])", re.IGNORECASE)),
    ("cover letter", re.compile(r"(?<![A-Za-z0-9])cover letter(?![A-Za-z0-9])", re.IGNORECASE)),
    ("transcript", token("transcript")),
    ("blank", token("blank")),
    ("scan", token("scan")),
]

COPY_INDEX = re.compile(r"\s*\((?!(?:19|20)\d{2}\))\d+\)")
COPY_OF = re.compile(r"\bcopy of\b", re.IGNORECASE)


def alnum_count(text: str) -> int:
    return sum(character.isalnum() for character in text)


def extract_pdf(path: Path) -> tuple[str, int]:
    """Return (status, alphanumeric count). Status is ok, empty, or failed:reason."""
    try:
        document = pymupdf.open(path)
    except Exception as exc:
        return f"failed:{type(exc).__name__}", 0

    try:
        if document.needs_pass:
            return "failed:encrypted", 0
        text = "".join(page.get_text() or "" for page in document)
    except Exception as exc:
        return f"failed:{type(exc).__name__}", 0
    finally:
        document.close()

    count = alnum_count(text)
    if count == 0:
        return "empty", 0
    return "ok", count


def extract_docx(path: Path) -> tuple[str, int]:
    try:
        document = docx.Document(path)
    except Exception as exc:
        return f"failed:{type(exc).__name__}", 0

    chunks: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)

    count = alnum_count("\n".join(chunks))
    if count == 0:
        return "empty", 0
    return "ok", count


def file_kind(path: Path) -> str:
    extension = path.suffix.lower()
    if extension == ".pdf":
        return "pdf"
    if extension == ".docx":
        return "docx"
    if extension in IMAGE_EXTENSIONS:
        return "image"
    return "other"


def normalized_stem(path: Path) -> str:
    """Filename with extension, copy indexes, and 'copy of' removed."""
    stem = COPY_OF.sub(" ", path.stem)
    stem = COPY_INDEX.sub(" ", stem)
    stem = re.sub(r"[\s_\-]+", " ", stem.casefold())
    return stem.strip()


def trailing_name(path: Path) -> str | None:
    """Text after the last ' - ', which is how many of these files name the person."""
    if " - " not in path.stem:
        return None
    tail = path.stem.rsplit(" - ", 1)[-1]
    tail = COPY_INDEX.sub(" ", tail)
    tail = re.sub(r"\s+", " ", tail).strip(" .-_")
    if len(tail) < 3:
        return None
    return tail.casefold()


def print_groups(title: str, groups: dict[str, list[str]]) -> None:
    repeated = {key: names for key, names in groups.items() if len(names) > 1}
    file_count = sum(len(names) for names in repeated.values())
    print(f"{title}: {len(repeated)} groups, {file_count} files")
    for key in sorted(repeated, key=lambda item: (-len(repeated[item]), item)):
        names = repeated[key]
        print(f"  [{len(names)}] {key}")
        for name in sorted(names):
            print(f"      {name}")


def inspect(cvs_dir: Path) -> int:
    files = sorted(path for path in cvs_dir.rglob("*") if path.is_file())
    counts = {"pdf": 0, "docx": 0, "image": 0, "other": 0}
    extractable: list[str] = []
    needs_ocr: list[str] = []
    empty: list[str] = []
    failed: list[tuple[str, str]] = []
    stem_groups: dict[str, list[str]] = defaultdict(list)
    trailing_groups: dict[str, list[str]] = defaultdict(list)
    version_hits: dict[str, list[str]] = defaultdict(list)
    non_cv_hits: dict[str, list[str]] = defaultdict(list)

    for path in files:
        kind = file_kind(path)
        counts[kind] += 1
        relative = path.relative_to(cvs_dir).as_posix()

        if kind == "pdf":
            status, count = extract_pdf(path)
        elif kind == "docx":
            status, count = extract_docx(path)
        elif kind == "image":
            status, count = "image", 0
        else:
            status, count = "failed:unsupported", 0

        if status == "ok" and count >= MIN_ALNUM:
            extractable.append(relative)
        elif status == "empty":
            empty.append(relative)
            if kind == "pdf":
                needs_ocr.append(relative)
        elif status == "ok":
            # Opened, but the text layer is too thin to be a real CV.
            needs_ocr.append(relative)
        elif status == "image":
            needs_ocr.append(relative)
        else:
            reason = status.removeprefix("failed:")
            failed.append((relative, reason))

        stem_key = normalized_stem(path)
        if stem_key:
            stem_groups[stem_key].append(relative)

        person = trailing_name(path)
        if person:
            trailing_groups[person].append(relative)

        filename = path.name
        for label, pattern in VERSION_PATTERNS:
            if pattern.search(filename):
                version_hits[label].append(relative)
        for label, pattern in NON_CV_PATTERNS:
            if pattern.search(filename):
                non_cv_hits[label].append(relative)

    version_files = sorted({name for names in version_hits.values() for name in names})

    print(f"Folder: {cvs_dir}")
    print(
        f"Local text check only (PyMuPDF, python-docx). "
        f"A file counts as extractable at {MIN_ALNUM}+ letters or digits. "
        "No text is printed or sent to a model."
    )
    print()
    print(f"Total files: {len(files)}")
    print(f"PDF: {counts['pdf']}")
    print(f"DOCX: {counts['docx']}")
    print(f"Images: {counts['image']}")
    print(f"Other: {counts['other']}")
    print()
    print(f"Files with extractable text: {len(extractable)}")
    print(f"Files requiring OCR: {len(needs_ocr)}")
    print(f"Files with empty extraction: {len(empty)}")
    print(f"Files that could not be opened: {len(failed)}")
    if failed:
        for name, reason in failed:
            print(f"  {name} ({reason})")
    if empty:
        print("Empty extraction:")
        for name in empty:
            print(f"  {name}")
    image_names = [name for name in needs_ocr if Path(name).suffix.lower() in IMAGE_EXTENSIONS]
    if image_names:
        print("Images (no text layer):")
        for name in image_names:
            print(f"  {name}")
    sparse = [name for name in needs_ocr if name not in empty and name not in image_names]
    if sparse:
        print(f"Thin text layer (under {MIN_ALNUM} letters or digits), still counted as OCR:")
        for name in sparse:
            print(f"  {name}")
    print()
    print_groups(
        "Potential duplicate filenames (same name aside from extension, '(n)', and 'copy of')",
        stem_groups,
    )
    print()
    print_groups(
        "Potential duplicate filenames (same trailing name after ' - ')",
        trailing_groups,
    )
    print()
    print(f"Potential version/update filenames: {len(version_files)}")
    for label, _pattern in VERSION_PATTERNS:
        names = version_hits.get(label, [])
        if not names:
            continue
        print(f"  {label}: {len(names)}")
        for name in names:
            print(f"      {name}")
    print()
    non_cv_files = sorted({name for names in non_cv_hits.values() for name in names})
    print(f"Filenames that do not look like a CV: {len(non_cv_files)}")
    for name in non_cv_files:
        print(f"  {name}")

    accounted = len(extractable) + len(needs_ocr) + len(failed)
    if accounted != len(files):
        print()
        print(f"Count mismatch: {accounted} classified, {len(files)} files on disk.")
        return 1
    return 0


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    parser = argparse.ArgumentParser(description="Inspect the local CV folder.")
    parser.add_argument(
        "cvs_dir",
        nargs="?",
        type=Path,
        help="Folder of CV files. Defaults to ../cvs next to this script.",
    )
    args = parser.parse_args()
    cvs_dir = args.cvs_dir or Path(__file__).resolve().parent.parent / "cvs"
    if not cvs_dir.is_dir():
        print(f"CV folder not found: {cvs_dir}", file=sys.stderr)
        sys.exit(1)
    sys.exit(inspect(cvs_dir))


if __name__ == "__main__":
    main()
